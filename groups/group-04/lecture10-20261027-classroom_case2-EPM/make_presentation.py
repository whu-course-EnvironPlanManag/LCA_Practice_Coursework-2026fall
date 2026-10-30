from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
FIGURES = ROOT / "figures_output"
OUTPUT = ROOT / "lecture10_classroom_presentation.pptx"
FONT = "Microsoft YaHei"
GREEN, GREEN2, ORANGE, CYAN = "173E36", "2F6B59", "D96B3B", "4D9A9A"
YELLOW, INK, MUTED = "E5B94F", "202A28", "65716D"
PAPER, WHITE, LINE, RED = "F4F5F0", "FFFFFF", "D8DED8", "B85645"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
prs.core_properties.title = "第十讲课堂汇报：绿色甲醇园区环境规划"
prs.core_properties.subject = "MFA-LCA-TEA-Policy 综合决策"


def color(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill, border=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if border:
        shape.line.color.rgb = color(border)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    return shape


def text(slide, x, y, w, h, value, size=14, fill=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for i, line in enumerate(str(value).split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(4)
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color(fill)
    return shape


def line(slide, x1, y1, x2, y2, fill=LINE, width=1):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color(fill)
    shape.line.width = Pt(width)
    return shape


def slide(title, section, number):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color(PAPER)
    text(s, .55, .22, 9, .25, section.upper(), 9, GREEN2, True)
    text(s, .55, .51, 12, .5, title, 24, GREEN, True)
    line(s, .55, 1.12, 12.78, 1.12)
    text(s, .55, 7.15, 8, .18, "环境规划与管理  /  第十讲课堂汇报", 8, MUTED)
    text(s, 12.1, 7.15, .55, .18, f"{number:02d}", 9, MUTED, True, PP_ALIGN.RIGHT)
    return s


def metric(s, x, y, number, caption, accent=ORANGE, unit=""):
    text(s, x, y, 2.45, .48, number, 26, accent, True)
    text(s, x, y + .52, 2.5, .25, unit, 9, MUTED, True)
    text(s, x, y + .82, 2.55, .42, caption, 11, INK, True)


def fit_picture(s, path, x, y, max_w, max_h):
    with Image.open(path) as image:
        aspect = image.width / image.height
    width, height = max_w, max_w / aspect
    if height > max_h:
        width, height = max_h * aspect, max_h
    x += (max_w - width) / 2
    y += (max_h - height) / 2
    s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(width), height=Inches(height))


# 01 Cover
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = color(GREEN)
box(s, 0, 0, .18, 7.5, ORANGE)
text(s, .82, .65, 6.5, .3, "ENVIRONMENTAL PLANNING  /  LECTURE 10", 10, "B8D3C5", True)
text(s, .82, 1.35, 7, 1.55, "绿色甲醇示范线\n环境规划与管理", 32, WHITE, True)
text(s, .86, 3.25, 6.5, .48, "碳捕集 × 绿氢 × 多目标决策", 19, "DDE8E1", True)
line(s, .86, 4.12, 7.35, 4.13, "6E9284")
text(s, .86, 4.42, 6.3, .72, "以 10 万吨/年绿色甲醇项目为对象，\n连接物质流、生命周期、经济性与政策约束。", 14, "DDE8E1")
text(s, .86, 6.68, 6.8, .25, "《环境规划与管理》  |  2026 秋季学期", 10, "B8D3C5")
box(s, 8.15, .78, 4.42, 5.92, "214C42")
text(s, 8.55, 1.12, 3.65, .25, "规划设计基准  /  ANCHOR B", 10, "B8D3C5", True)
metric(s, 8.55, 1.72, "100,000", "甲醇产能", YELLOW, "t/a")
line(s, 8.55, 3.02, 12.1, 3.03, "6E9284")
metric(s, 8.55, 3.25, "−1,066", "净碳足迹", "8CC9B5", "kg CO₂e/t")
line(s, 8.55, 4.58, 12.1, 4.59, "6E9284")
metric(s, 8.55, 4.82, "2,309.5", "调整后 LCOM", "F09969", "元/t")
text(s, 8.55, 6.42, 3.65, .22, "模型情景值，非项目实测或审批结论", 8, "C5D8D0")

# 02 Planning question
s = slide("为什么需要系统级规划？", "01  /  Planning question", 2)
text(s, .7, 1.48, 4.1, .4, "目标不是单独增加一台电解槽", 19, GREEN, True)
text(s, .7, 2.05, 4.05, 1.0, "项目同时受连续化工生产、波动性可再生电力、用水边界与产品碳足迹约束。", 15)
box(s, .7, 3.48, 4.0, 2.55, WHITE, LINE)
text(s, .96, 3.75, 3.4, .25, "规划对象", 11, GREEN2, True)
text(s, .96, 4.18, 3.5, 1.35, "大型煤化工园区\nCO₂ 捕集 + PEM 制氢 + 甲醇合成\n设计产能：100,000 t/a", 15, INK, True)
text(s, 5.4, 1.5, 6.7, .35, "四类分析放进同一决策框架", 18, GREEN, True)
for i, (tag, desc, c) in enumerate([
    ("MFA", "循环反应与物料守恒", GREEN2), ("LCA", "电力结构与产品 GWP", CYAN),
    ("TEA", "成本、资本年金与 MAC", ORANGE), ("POLICY", "认证、碳、水与污染约束", "7B7B47")]):
    y = 2.08 + i * .92
    box(s, 5.4, y, 1.42, .6, c)
    text(s, 5.48, y + .14, 1.25, .22, tag, 11, WHITE, True, PP_ALIGN.CENTER)
    text(s, 7.15, y + .11, 4.7, .32, desc, 14, INK, True)
text(s, 5.4, 6.05, 6.5, .4, "决策变量：绿电比例 f_green × 循环气驰放比 r_purge", 12, GREEN2, True)

# 03 Process diagram
s = slide("从捕集到产品：物质与能量如何闭合？", "02  /  System boundary", 3)
img = FIGURES / "fig1_mfa_process_flow.png"
if img.exists():
    fit_picture(s, img, .55, 1.34, 12.23, 5.5)
text(s, .62, 6.91, 11.6, .18, "锚点 B 流程图；循环气为内部物流，不重复计入外部输入。", 8, MUTED)

# 04 Power and operation
s = slide("绿电比例决定碳结果，供电稳定性决定装置能否运行", "03  /  Power & operations", 4)
box(s, .68, 1.48, 5.72, 4.82, GREEN)
text(s, 1.04, 1.82, 4.9, .3, "课堂手册：煤电 / 光伏混合供电临界值", 12, "B8D3C5", True)
text(s, 1.04, 2.48, 4.9, .65, "x ≥ 53.4%", 34, WHITE, True)
text(s, 1.04, 3.42, 4.95, .95, "GWP_net(x) = 10,826 × (0.75 − 0.71x) − 1,423\n需低于煤制甲醇基准 +2,600 kg CO₂e/t", 13, "DDE8E1")
line(s, 1.04, 4.75, 5.84, 4.76, "6E9284")
text(s, 1.04, 5.02, 4.95, .9, "该临界值来自课堂手册参数，不等同于任务 3 锚点 B 的电力因子与优化模型。", 10, "DDE8E1")
text(s, 7.02, 1.58, 5.05, .6, "化工连续运行与光伏间歇出力存在时间错配", 17, GREEN, True)
text(s, 7.05, 2.3, 5.0, 2.75, "8,000 h/a  固定床合成的设计运行小时要求\n\n15%  无储能离网光伏供电时的设备利用率假设\n\n6.67×  利用率下降导致单位产品设备年金摊销上升", 14)
box(s, 7.03, 5.42, 5.05, .84, "E5ECE5")
text(s, 7.28, 5.61, 4.6, .43, "绿电直供 + 柔性电解 + 储氢缓冲 + 合成端稳态运行", 11, GREEN, True)

# 05 Carbon scenarios
s = slide("高碳电力会抵消电制甲醇的减排收益", "04  /  Lifecycle impact", 5)
text(s, .7, 1.4, 11.6, .28, "同一任务 3 循环工况下仅改变电力结构；煤制甲醇为独立行业对照基准。", 11, MUTED)
left, width = 3.5, 8.25
zero = left + width * 1500 / 6500
line(s, zero, 2.05, zero, 5.85, "9BA6A0", 1)
rows = [
    ("煤制甲醇基准", 2600.0, "行业对照", "7C8580"),
    ("非绿电对照", 4560.99, "0% 绿电", RED),
    ("多能互补过渡", 903.39, "65% 绿电", ORANGE),
    ("锚点 B 推荐", -1066.08, "100% 绿电", GREEN2),
]
for i, (name, value, note, c) in enumerate(rows):
    y = 2.2 + i * .85
    text(s, .75, y + .04, 2.45, .25, name, 12, INK, True)
    text(s, .75, y + .34, 2.45, .18, note, 9, MUTED)
    xv = left + width * (value + 1500) / 6500
    if value >= 0:
        box(s, zero, y, xv-zero, .34, c)
        text(s, xv + .1, y + .02, 1.0, .25, f"{value:+,.0f}", 9, INK, True)
    else:
        box(s, xv, y, zero-xv, .34, c)
        text(s, max(left, xv-.92), y + .02, .9, .25, f"{value:+,.0f}", 9, INK, True)
for tick in [-1000, 0, 1000, 2000, 3000, 4000, 5000]:
    x = left + width * (tick+1500)/6500
    line(s, x, 5.92, x, 6.02, "9BA6A0", .6)
    text(s, x-.3, 6.08, .68, .2, f"{tick:+,}", 8, MUTED, False, PP_ALIGN.CENTER)
text(s, 8.45, 6.37, 3.6, .24, "GWP · kg CO₂e / t 甲醇", 10, GREEN, True, PP_ALIGN.RIGHT)
text(s, .75, 6.7, 11.5, .18, "模型结果受排放因子、捕集碳信用和系统边界影响，不替代园区实测清单或第三方核查。", 8, MUTED)

# 06 Pareto anchors
s = slide("锚点 B：在成本与净碳足迹之间选取折中方案", "05  /  Multi-objective decision", 6)
img = FIGURES / "fig2_pareto_decision_frontier.png"
if img.exists():
    fit_picture(s, img, .62, 1.44, 7.12, 4.62)
text(s, .7, 6.15, 6.9, .38, "3,000 组候选 → 政策筛除 → 非支配前沿 → 理想点距离选 B", 9, MUTED)
text(s, 8.0, 1.52, 4.6, .3, "三大决策锚点", 15, GREEN, True)
cols = [("方案", 1.2), ("驰放比", .82), ("GWP", .8), ("LCOM", 1.05)]
x = 8.0
for name, w in cols:
    box(s, x, 2.0, w, .4, GREEN)
    text(s, x+.03, 2.09, w-.06, .2, name, 9, WHITE, True, PP_ALIGN.CENTER)
    x += w
for i, row in enumerate([
    ("A 成本", "0.84%", "−1,038", "2,235.9"),
    ("B 拐点", "2.16%", "−1,066", "2,309.5"),
    ("C 深度减碳", "3.50%", "−1,092", "2,401.6")]):
    y = 2.42 + i*.63
    box(s, 8.0, y, 3.87, .55, "E0EBE3" if i == 1 else WHITE, LINE)
    x = 8.0
    for (_, w), val in zip(cols, row):
        text(s, x+.03, y+.15, w-.06, .22, val, 9, GREEN if i == 1 else INK, i == 1, PP_ALIGN.CENTER)
        x += w
box(s, 8.0, 4.55, 4.45, 1.2, GREEN)
text(s, 8.22, 4.78, 4.05, .74, "B 对 A：GWP 再降约 28 kg/t\nB 对 C：成本低约 92 元/t，碳利用率更高", 11, WHITE, True)
text(s, 8.0, 6.12, 4.5, .48, "锚点 B 碳利用率 93.92%；各数值均为模型输出。", 8, MUTED)

# 07 Engineering scale
s = slide("锚点 B 的年通量：10 万吨产能对应的系统规模", "06  /  Engineering scale", 7)
for x, val, caption, unit, c in [
    (.78,"100,000","甲醇产品","t/a",GREEN2), (3.85,"146,240","新鲜捕集 CO₂","t/a",CYAN),
    (6.92,"20,097","电解绿氢","t/a",ORANGE), (9.99,"1.093","综合用电","TWh/a",GREEN)]:
    box(s, x, 1.6, 2.55, 1.62, WHITE, LINE)
    metric(s, x+.2, 1.82, val, caption, c, unit)
text(s, .82, 3.7, 5.5, .3, "水与循环物流", 16, GREEN, True)
line(s, .9, 4.32, 12.28, 4.33, LINE, 1)
for x, val, caption, c in [
    (1.0,"138,500 t/a","新鲜净补水",CYAN), (4.0,"56,223 t/a","精馏水回用",GREEN2),
    (7.0,"458,569 t/a","未反应气回流",ORANGE), (10.0,"2.157%","驰放比",RED)]:
    box(s, x, 4.12, .16, .44, c)
    text(s, x, 4.66, 2.55, .33, val, 15, c, True)
    text(s, x, 5.08, 2.55, .3, caption, 11, INK, True)
box(s, .82, 5.86, 11.6, .57, "E8EDE7")
text(s, 1.02, 6.03, 11.1, .22, "任务 1 默认 1.0% 驰放比与锚点 B 不同；正式设计前须统一物料衡算边界。", 9, GREEN, True)

# 08 Economics
s = slide("经济性结论必须区分模型成本与工程概算", "07  /  Techno-economics", 8)
box(s, .72, 1.5, 4.05, 4.9, GREEN)
text(s, 1.03, 1.83, 3.35, .28, "锚点 B  ·  经济基准", 11, "B8D3C5", True)
metric(s, 1.03, 2.35, "2,309.5", "政策调整后 LCOM", "F09969", "元/t")
line(s, 1.03, 3.72, 4.35, 3.73, "6E9284")
text(s, 1.03, 4.0, 3.35, .28, "模型隐含总 CAPEX", 12, "DDE8E1", True)
text(s, 1.03, 4.46, 3.35, .48, "约 14.88 亿元", 23, YELLOW, True)
text(s, 1.03, 5.28, 3.35, .83, "由资本年金反推，非项目审定概算；工程分项投资与 OPEX 尚未输出。", 10, "DDE8E1")
text(s, 5.42, 1.58, 6.3, .3, "LCOM 构成（元/t）", 17, GREEN, True)
for i, (name, val, c) in enumerate([
    ("综合电力",1966.8,ORANGE), ("资本年金",1515.2,GREEN2), ("捕集 CO₂",380.2,CYAN),
    ("水与运维",350.0,"7B7B47"), ("副产氧收益",-321.5,RED)]):
    y = 2.16+i*.66
    text(s, 5.44, y+.04, 1.38, .23, name, 10, INK, True)
    box(s, 7.0, y, max(abs(val)/2000*3.5,.08), .32, c)
    text(s, 10.72, y+.03, 1.25, .24, f"{val:+,.0f}", 10, INK, True)
box(s, 5.42, 5.62, 6.55, .62, "E8EDE7")
text(s, 5.66, 5.78, 6.0, .28, "WACC 8% · 20 年 · CRF 0.10185 · 电价每 +0.01 元/kWh，LCOM 约 +109 元/t", 9, GREEN, True)
text(s, 5.46, 6.44, 6.45, .26, "补贴、绿色溢价、CCER 与氧气销售均需现实核证或合同支持。", 8, MUTED)

# 09 Engineering plan
s = slide("工程落位按“源—网—荷—储—化”组织", "08  /  Implementation", 9)
stages = [("碳源","捕集与压缩","邻近热电点源",CYAN),("电力","绿电接入","资源与并网校核",YELLOW),
          ("制氢","PEM + 储氢","独立安全边界",ORANGE),("合成","反应与循环","连续稳态运行",GREEN2),
          ("产品","精馏与回水","产品外送 / 水回用",GREEN)]
for i, (name, detail, note, c) in enumerate(stages):
    x=.73+i*2.48
    box(s,x,1.62,2.05,1.08,c)
    text(s,x+.12,1.77,1.8,.26,name,15,WHITE,True,PP_ALIGN.CENTER)
    text(s,x+.12,2.16,1.8,.22,detail,10,WHITE,True,PP_ALIGN.CENTER)
    text(s,x+.06,2.84,1.95,.28,note,9,MUTED,False,PP_ALIGN.CENTER)
text(s,.76,3.55,11.7,.32,"五项重点工程包  /  CAPEX、OPEX 分项待可研与 FEED 核定",15,GREEN,True)
for i,(num,title,detail) in enumerate([
    ("01","CO₂ 捕集及压缩","14.624 万吨 CO₂/a"),("02","PEM 制氢及储氢","2.010 万吨 H₂/a"),
    ("03","甲醇合成与循环","10 万吨产品/a"),("04","精馏及水回用","5.62 万吨回用水/a"),
    ("05","管廊与能碳 MRV","物流计量 / 绿电溯源")]):
    y=4.04+i*.43
    text(s,.85,y,.43,.23,num,10,ORANGE,True)
    text(s,1.42,y,3.3,.23,title,11,INK,True)
    text(s,5.05,y,5.85,.23,detail,10,MUTED)
text(s,.84,6.4,11.7,.34,"总 CAPEX 14.88 亿元为模型量级估算；氢区间距与应急池容量须专项风险评价确认。",8,MUTED)

# 10 Governance and recommendation
s = slide("近期主推锚点 B，但以工程核实和合规条件为前置", "09  /  Governance & recommendation", 10)
text(s,.78,1.48,5.45,.32,"实施前四项核验",17,GREEN,True)
for i,(head,detail) in enumerate([
    ("电力","100% 绿电的逐时供给、溯源与并网能力"),
    ("物料","统一任务 1 / 任务 3 的进料及碳利用率边界"),
    ("水资源","取水许可、枯水期保障与回用水质"),
    ("安全环境","氢气 / CO₂ 风险、烟气基线与胺液废液控制")]):
    y=2.02+i*.82
    box(s,.8,y,1.28,.47,GREEN2 if i<2 else ORANGE)
    text(s,.9,y+.11,1.08,.22,head,11,WHITE,True,PP_ALIGN.CENTER)
    text(s,2.27,y+.07,4.0,.34,detail,10,INK,True)
box(s,6.95,1.5,5.6,3.42,GREEN)
text(s,7.34,1.84,4.8,.28,"结论  /  近期设计建议",12,"B8D3C5",True)
text(s,7.34,2.3,4.7,1.48,"采用锚点 B 作为方案比选基准：\n100% 绿电 · 2.157% 驰放比\n−1,066 kg CO₂e/t · 2,309.5 元/t",17,WHITE,True)
text(s,7.34,4.17,4.7,.4,"目标折中不代表已具备项目核准条件。",10,"DDE8E1")
text(s,.82,5.68,11.55,.32,"管理闭环：计量与 MRV  →  合规审查  →  风险演练  →  PDCA 动态修编",14,GREEN,True)
line(s,.84,6.18,12.35,6.19,LINE)
text(s,.85,6.4,11.7,.38,"核心信息：电力来源决定减碳成效；物料与政策边界决定模型可信度；工程可行性须由现场数据和专项设计验证。",10,INK,True)

prs.save(OUTPUT)
print(f"Created {OUTPUT}")
print(f"Slides: {len(prs.slides)}")